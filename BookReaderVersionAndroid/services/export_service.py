import csv
from pathlib import Path

from models.book import Book


def export_books_to_csv(books: list[Book], file_path: str | Path) -> Path:
    path = Path(file_path)
    if path.suffix.lower() != ".csv":
        path = path.with_suffix(".csv")

    path.parent.mkdir(parents=True, exist_ok=True)

    with open(path, "w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            ("Название", "Автор", "Всего страниц", "Текущая страница", "Прогресс (%)", "Статус")
        )
        for book in books:
            status = "Завершена" if book.progress >= 100 else "В процессе"
            writer.writerow(
                (
                    book.title,
                    book.author,
                    book.total_pages,
                    book.current_page,
                    book.progress,
                    status,
                )
            )

    return path


def export_books_to_excel(books: list[Book], file_path: str | Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise ImportError(
            "Для Excel установите openpyxl: pip install openpyxl"
        ) from error

    path = Path(file_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        path = path.with_suffix(".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Библиотека"

    headers = (
        "Название",
        "Автор",
        "Всего страниц",
        "Текущая страница",
        "Прогресс (%)",
        "Статус",
    )
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E66F5", end_color="1E66F5", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_index, book in enumerate(books, start=2):
        status = "Завершена" if book.progress >= 100 else "В процессе"
        row_values = (
            book.title,
            book.author,
            book.total_pages,
            book.current_page,
            book.progress,
            status,
        )
        for column, value in enumerate(row_values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    column_widths = (30, 25, 16, 18, 14, 14)
    for column, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
