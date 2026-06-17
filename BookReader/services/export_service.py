from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.work import WORK_STATUS_LABELS


BOOK_HEADERS = (
    "Название",
    "Автор",
    "Всего страниц",
    "Текущая страница",
    "Прогресс (%)",
    "Статус",
    "Описание",
    "Ссылка на маркетплейс",
    "Мнение",
    "Произведений",
    "Прочитано произведений",
)

WORK_HEADERS = (
    "Книга",
    "Автор",
    "Произведение",
    "Страниц",
    "Прочитано",
    "Прогресс (%)",
    "Статус",
)


class ExportService:

    @staticmethod
    def export_books_to_excel(books, file_path):
        path = Path(file_path)

        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            path = path.with_suffix(".xlsx")

        workbook = Workbook()
        books_sheet = workbook.active
        books_sheet.title = "Библиотека"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E66F5", end_color="1E66F5", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        wrap_alignment = Alignment(wrap_text=True, vertical="top")

        for column, header in enumerate(BOOK_HEADERS, start=1):
            cell = books_sheet.cell(row=1, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        for row_index, book in enumerate(books, start=2):
            status = "Завершена" if book.is_finished else "В процессе"
            works_count = len(book.works)
            finished_works = sum(1 for work in book.active_works if work.is_finished)
            row_values = (
                book.title,
                book.author,
                book.total_pages,
                book.effective_current_page,
                book.progress,
                status,
                book.description,
                book.marketplace_url,
                book.review,
                works_count,
                finished_works,
            )

            for column, value in enumerate(row_values, start=1):
                cell = books_sheet.cell(row=row_index, column=column, value=value)
                if column in {7, 9}:
                    cell.alignment = wrap_alignment
                if column == 8 and value:
                    cell.hyperlink = value
                    cell.font = Font(color="0563C1", underline="single")

        book_column_widths = (30, 25, 16, 18, 14, 14, 40, 35, 40, 16, 22)
        for column, width in enumerate(book_column_widths, start=1):
            books_sheet.column_dimensions[get_column_letter(column)].width = width

        books_sheet.freeze_panes = "A2"

        works_sheet = workbook.create_sheet("Произведения")
        for column, header in enumerate(WORK_HEADERS, start=1):
            cell = works_sheet.cell(row=1, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        work_row = 2
        for book in books:
            for work in book.works:
                status_label = WORK_STATUS_LABELS.get(work.status, work.status)
                row_values = (
                    book.title,
                    book.author,
                    work.title,
                    work.total_pages,
                    work.current_page,
                    work.progress,
                    status_label,
                )
                for column, value in enumerate(row_values, start=1):
                    works_sheet.cell(row=work_row, column=column, value=value)
                work_row += 1

        work_column_widths = (30, 25, 35, 12, 14, 14, 18)
        for column, width in enumerate(work_column_widths, start=1):
            works_sheet.column_dimensions[get_column_letter(column)].width = width

        works_sheet.freeze_panes = "A2"

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)

        return path
